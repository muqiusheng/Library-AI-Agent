import hashlib
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from openpyxl import load_workbook
from psycopg.types.json import Jsonb
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import (
    DataVersion,
    DataVersionStatus,
    ImportBatch,
    ImportStatus,
)
from app.pipelines.clean import (
    clean_text,
    map_book_status,
    map_circulation_status,
    normalize_for_match,
    normalize_isbn,
    parse_date,
    parse_decimal,
    parse_pages,
    parse_year,
)

CLEANING_VERSION = "bookpath-clean-v1"
REQUIRED_COLUMNS = {"序号", "题名", "索书号", "馆藏地", "所在地", "书刊状态", "流通状态", "ISBN"}
BATCH_SIZE = 2_000
BOOK_COLUMNS = (
    "id",
    "data_version_id",
    "isbn10",
    "isbn13",
    "title",
    "normalized_title",
    "responsible_person",
    "responsible_note",
    "publisher",
    "publication_year",
    "description",
    "pages",
    "price",
    "volume_issue",
    "source_fingerprint",
    "created_at",
    "updated_at",
)
HOLDING_COLUMNS = (
    "id",
    "book_id",
    "source_row_no",
    "call_number",
    "collection_location",
    "current_location",
    "shelf_number",
    "book_status",
    "circulation_status",
    "acquired_at",
    "updated_source_at",
    "raw_payload",
    "created_at",
    "updated_at",
)
ISSUE_COLUMNS = (
    "id",
    "batch_id",
    "row_no",
    "field_name",
    "issue_type",
    "raw_value",
    "description",
    "status",
    "created_at",
    "updated_at",
)


@dataclass(frozen=True)
class ImportReport:
    batch_id: str
    data_version_id: str
    duplicate_file: bool
    raw_rows: int
    valid_rows: int
    skipped_rows: int
    books: int
    holdings: int
    isbn_valid: int
    isbn_missing: int
    isbn_invalid: int
    issues: int
    status_distribution: dict[str, int]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _fingerprint(parts: list[Any]) -> str:
    text = "|".join(clean_text(part) or "" for part in parts)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _aggregate_key(row: dict[str, Any], isbn10: str | None, isbn13: str | None) -> str:
    if isbn13:
        return f"isbn13:{isbn13}"
    if isbn10:
        return f"isbn10:{isbn10}"
    return "meta:" + _fingerprint(
        [row.get("题名"), row.get("责任者"), row.get("出版社"), row.get("出版年"), row.get("卷期")]
    )


def _chunks[T](items: list[T], size: int = BATCH_SIZE):
    for start in range(0, len(items), size):
        yield items[start : start + size]


class ExcelLibraryImporter:
    def __init__(self, session: Session) -> None:
        self.session = session

    def import_file(self, file_path: str | Path, *, publish: bool = True) -> ImportReport:
        path = Path(file_path).resolve()
        if path.suffix.lower() != ".xlsx" or not path.is_file():
            raise ValueError("only an existing .xlsx file can be imported")
        file_hash = file_sha256(path)
        existing = self.session.scalar(
            select(ImportBatch).where(ImportBatch.file_hash == file_hash)
        )
        if existing and existing.status == ImportStatus.SUCCEEDED:
            stats = existing.statistics
            return ImportReport(
                batch_id=str(existing.id),
                data_version_id=str(existing.data_version_id),
                duplicate_file=True,
                **{
                    key: stats[key]
                    for key in ImportReport.__dataclass_fields__
                    if key not in {"batch_id", "data_version_id", "duplicate_file"}
                },
            )

        batch = existing or ImportBatch(
            file_name=path.name,
            file_hash=file_hash,
            stored_path=str(path),
            status=ImportStatus.PENDING,
            stage="validating",
            progress=0,
            statistics={},
        )
        if existing is None:
            self.session.add(batch)
        batch.status = ImportStatus.RUNNING
        batch.started_at = datetime.now(UTC)
        self.session.commit()

        try:
            return self._run(path, batch, publish=publish)
        except Exception as exc:
            self.session.rollback()
            failed = self.session.get(ImportBatch, batch.id)
            if failed:
                failed.status = ImportStatus.FAILED
                failed.error_summary = f"{exc.__class__.__name__}: {exc}"[:2000]
                failed.finished_at = datetime.now(UTC)
                self.session.commit()
            raise

    def _run(self, path: Path, batch: ImportBatch, *, publish: bool) -> ImportReport:
        version = DataVersion(
            name=f"library-{datetime.now(UTC):%Y%m%d-%H%M%S}-{batch.file_hash[:8]}",
            status=DataVersionStatus.PROCESSING,
        )
        self.session.add(version)
        self.session.flush()
        batch.data_version_id = version.id
        batch.stage = "cleaning"

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows)
        headers = [
            clean_text(value) or f"column_{index}" for index, value in enumerate(raw_headers)
        ]
        missing = REQUIRED_COLUMNS - set(headers)
        if missing:
            raise ValueError(f"missing required columns: {', '.join(sorted(missing))}")

        book_rows: dict[str, dict[str, Any]] = {}
        holding_rows: list[dict[str, Any]] = []
        issue_rows: list[dict[str, Any]] = []
        isbn_counts = Counter(valid=0, missing=0, invalid=0)
        status_counts: Counter[str] = Counter()
        raw_count = valid_count = skipped_count = 0

        for excel_row_no, values in enumerate(rows, start=2):
            raw_count += 1
            row = {
                header: clean_text(value) for header, value in zip(headers, values, strict=False)
            }
            title = row.get("题名")
            if not title:
                skipped_count += 1
                self._issue(
                    issue_rows,
                    batch.id,
                    excel_row_no,
                    "题名",
                    "missing_required",
                    row.get("题名"),
                    "题名为空",
                )
                continue
            isbn10, isbn13, invalid_isbn = normalize_isbn(row.get("ISBN"))
            if isbn10 or isbn13:
                isbn_counts["valid"] += 1
            elif invalid_isbn:
                isbn_counts["invalid"] += 1
                self._issue(
                    issue_rows,
                    batch.id,
                    excel_row_no,
                    "ISBN",
                    "invalid_isbn",
                    invalid_isbn,
                    "ISBN 校验位无效",
                )
            else:
                isbn_counts["missing"] += 1

            key = _aggregate_key(row, isbn10, isbn13)
            book_id = book_rows.get(key, {}).get("id") or uuid4()
            publication_year = parse_year(row.get("出版年"))
            source_fingerprint = hashlib.sha256(f"{CLEANING_VERSION}|{key}".encode()).hexdigest()
            if key not in book_rows:
                book_rows[key] = {
                    "id": book_id,
                    "data_version_id": version.id,
                    "isbn10": isbn10,
                    "isbn13": isbn13,
                    "title": title,
                    "normalized_title": normalize_for_match(title),
                    "responsible_person": row.get("责任者"),
                    "responsible_note": row.get("责任者附注"),
                    "publisher": row.get("出版社"),
                    "publication_year": publication_year,
                    "description": row.get("简介"),
                    "pages": parse_pages(row.get("页码")),
                    "price": parse_decimal(row.get("价格")),
                    "volume_issue": row.get("卷期"),
                    "source_fingerprint": source_fingerprint,
                    "created_at": datetime.now(UTC),
                    "updated_at": datetime.now(UTC),
                }
            book_status = map_book_status(row.get("书刊状态"), row.get("流通状态"))
            circulation_status = map_circulation_status(row.get("书刊状态"), row.get("流通状态"))
            status_counts[book_status] += 1
            source_row_no = int(float(row.get("序号") or excel_row_no - 1))
            holding_rows.append(
                {
                    "id": uuid4(),
                    "book_id": book_id,
                    "source_row_no": source_row_no,
                    "call_number": row.get("索书号"),
                    "collection_location": row.get("馆藏地"),
                    "current_location": row.get("所在地"),
                    "shelf_number": row.get("排架号"),
                    "book_status": book_status,
                    "circulation_status": circulation_status,
                    "acquired_at": parse_date(row.get("入藏时间")),
                    "updated_source_at": self._as_datetime(row.get("更新时间")),
                    "raw_payload": {
                        header: value for header, value in row.items() if value is not None
                    },
                    "created_at": datetime.now(UTC),
                    "updated_at": datetime.now(UTC),
                }
            )
            valid_count += 1
            if raw_count % 20_000 == 0:
                print(f"cleaned {raw_count}/{sheet.max_row - 1} rows", file=sys.stderr, flush=True)
        workbook.close()

        batch.stage = "deduplicating"
        print(f"copying {len(book_rows)} books", file=sys.stderr, flush=True)
        self._copy_rows("books", BOOK_COLUMNS, list(book_rows.values()))
        print(f"copying {len(holding_rows)} holdings", file=sys.stderr, flush=True)
        self._copy_rows("holdings", HOLDING_COLUMNS, holding_rows)
        if issue_rows:
            self._copy_rows("import_issues", ISSUE_COLUMNS, issue_rows)

        version.status = DataVersionStatus.PUBLISHED if publish else DataVersionStatus.DRAFT
        version.published_at = datetime.now(UTC) if publish else None
        stats = {
            "raw_rows": raw_count,
            "valid_rows": valid_count,
            "skipped_rows": skipped_count,
            "books": len(book_rows),
            "holdings": len(holding_rows),
            "isbn_valid": isbn_counts["valid"],
            "isbn_missing": isbn_counts["missing"],
            "isbn_invalid": isbn_counts["invalid"],
            "issues": len(issue_rows),
            "status_distribution": dict(status_counts),
        }
        batch.status = ImportStatus.SUCCEEDED
        batch.stage = "published" if publish else "ready_to_publish"
        batch.progress = 100
        batch.statistics = stats
        batch.finished_at = datetime.now(UTC)
        self.session.commit()
        return ImportReport(
            batch_id=str(batch.id),
            data_version_id=str(version.id),
            duplicate_file=False,
            **stats,
        )

    @staticmethod
    def _issue(
        rows: list[dict[str, Any]],
        batch_id: UUID,
        row_no: int,
        field: str,
        issue_type: str,
        raw: Any,
        description: str,
    ) -> None:
        rows.append(
            {
                "id": uuid4(),
                "batch_id": batch_id,
                "row_no": row_no,
                "field_name": field,
                "issue_type": issue_type,
                "raw_value": clean_text(raw),
                "description": description,
                "status": "open",
                "created_at": datetime.now(UTC),
                "updated_at": datetime.now(UTC),
            }
        )

    @staticmethod
    def _as_datetime(value: Any) -> datetime | None:
        parsed = parse_date(value)
        return datetime.combine(parsed, datetime.min.time(), tzinfo=UTC) if parsed else None

    def _copy_rows(
        self,
        table: str,
        columns: tuple[str, ...],
        rows: list[dict[str, Any]],
    ) -> None:
        if not rows:
            return
        driver_connection = self.session.connection().connection.driver_connection
        column_sql = ", ".join(columns)
        with driver_connection.cursor().copy(f"COPY {table} ({column_sql}) FROM STDIN") as copy:
            for row in rows:
                values = []
                for column in columns:
                    value = row.get(column)
                    values.append(Jsonb(value) if isinstance(value, (dict, list)) else value)
                copy.write_row(values)


def report_as_dict(report: ImportReport) -> dict[str, Any]:
    return asdict(report)
